from __future__ import annotations

import asyncio
import base64
import hashlib
import hmac
import io
import json
import os
import re
import time
import unicodedata
from contextlib import asynccontextmanager
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import psycopg
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from psycopg.rows import dict_row

ROOT = Path(__file__).resolve().parent
PUBLIC_DIR = ROOT / "public"
MIGRATIONS_DIR = ROOT / "database" / "migrations"
COOKIE_NAME = "prcs_hr_session"
SESSION_HOURS = 8
MAX_UPLOAD_BYTES = 16 * 1024 * 1024
DB_CONFIG = {
    "host": os.getenv("DB_HOST", "db"),
    "port": int(os.getenv("DB_PORT", "5432")),
    "dbname": os.getenv("DB_NAME", "prcs_hr"),
    "user": os.getenv("DB_USER", "prcs_hr_app"),
    "password": os.getenv("DB_PASSWORD", ""),
    "connect_timeout": 5,
}
SESSION_SECRET = os.getenv("SESSION_SECRET", "")
COOKIE_SECURE = os.getenv("COOKIE_SECURE", "true").lower() in ("1", "true", "yes")
LOGIN_FAILURES: dict[str, list[float]] = {}

if len(SESSION_SECRET) < 32:
    raise RuntimeError("Set a 32+ character SESSION_SECRET.")

def admin_setup_complete() -> bool:
    with connect() as conn:
        row = conn.execute("SELECT 1 FROM public.admin_account WHERE id = 1").fetchone()
    return row is not None

def get_admin_account():
    with connect() as conn:
        return conn.execute(
            "SELECT username, password_hash FROM public.admin_account WHERE id = 1"
        ).fetchone()

def password_hash(password: str) -> str:
    salt = os.urandom(16)
    digest = hashlib.scrypt(password.encode("utf-8"), salt=salt, n=2**14, r=8, p=1)
    return base64.b64encode(salt + digest).decode("ascii")

def verify_password(password: str, encoded_hash: str) -> bool:
    try:
        raw = base64.b64decode(encoded_hash, validate=True)
        salt, expected = raw[:16], raw[16:]
        actual = hashlib.scrypt(password.encode("utf-8"), salt=salt, n=2**14, r=8, p=1)
        return hmac.compare_digest(actual, expected)
    except (ValueError, TypeError):
        return False


def connect():
    return psycopg.connect(**DB_CONFIG, row_factory=dict_row)


def prepare_database() -> None:
    last_error: Exception | None = None
    for _ in range(40):
        try:
            with connect() as conn:
                for migration in sorted(MIGRATIONS_DIR.glob("*.sql")):
                    conn.execute(migration.read_text(encoding="utf-8"))
            return
        except Exception as exc:
            last_error = exc
            time.sleep(2)
    raise RuntimeError(f"Database is not ready: {last_error}")


@asynccontextmanager
async def lifespan(_: FastAPI):
    await asyncio.to_thread(prepare_database)
    yield


app = FastAPI(title="PRCS HR", docs_url=None, redoc_url=None, lifespan=lifespan)


def session_token(username: str, expires: int) -> str:
    payload = base64.urlsafe_b64encode(
        json.dumps({"sub": username, "exp": expires}, separators=(",", ":")).encode()
    ).decode().rstrip("=")
    signature = hmac.new(SESSION_SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()
    return f"{payload}.{signature}"


def authenticated_username(request: Request) -> str | None:
    token = request.cookies.get(COOKIE_NAME, "")
    try:
        payload, signature = token.rsplit(".", 1)
        expected = hmac.new(SESSION_SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(signature, expected):
            return None
        decoded = base64.urlsafe_b64decode(payload + "=" * (-len(payload) % 4))
        claims = json.loads(decoded)
        if int(claims["exp"]) <= int(time.time()):
            return None
        account = get_admin_account()
        if not account or not hmac.compare_digest(str(claims["sub"]), account["username"]):
            return None
        return str(claims["sub"])
    except (ValueError, KeyError, TypeError, json.JSONDecodeError):
        return None


def check_same_origin(request: Request) -> None:
    origin = request.headers.get("origin")
    if not origin:
        return
    host = request.headers.get("x-forwarded-host") or request.headers.get("host", "")
    scheme = request.headers.get("x-forwarded-proto", request.url.scheme).split(",")[0].strip()
    expected = f"{scheme}://{host}".rstrip("/")
    if not hmac.compare_digest(origin.rstrip("/"), expected):
        raise HTTPException(status_code=403, detail="مصدر الطلب غير مسموح.")


@app.middleware("http")
async def protect_application(request: Request, call_next):
    path = request.url.path
    open_paths = {"/login.html", "/api/login", "/api/session", "/api/setup-status", "/api/setup-admin", "/api/healthz"}
    if path not in open_paths:
        if path.startswith("/api/") and path not in {"/api/setup-status", "/api/setup-admin"} and authenticated_username(request) is None:
            return JSONResponse(status_code=401, content={"detail": "يلزم تسجيل الدخول."})
        if not path.startswith("/api/") and authenticated_username(request) is None:
            return RedirectResponse("/login.html", status_code=303)
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "same-origin"
    response.headers["Cache-Control"] = "no-store"
    return response


@app.get("/api/healthz")
def healthz():
    try:
        with connect() as conn:
            conn.execute("SELECT 1")
        return {"status": "ok"}
    except Exception:
        return JSONResponse(status_code=503, content={"status": "database_unavailable"})


@app.get("/api/setup-status")
def setup_status():
    return {"setup_required": not admin_setup_complete()}


@app.post("/api/setup-admin")
async def setup_admin(request: Request):
    check_same_origin(request)
    if admin_setup_complete():
        raise HTTPException(status_code=409, detail="تم إنشاء حساب المدير مسبقاً.")
    try:
        body = await request.json()
    except Exception:
        raise HTTPException(status_code=400, detail="بيانات إنشاء الحساب غير صالحة.")
    username = str(body.get("username", "")).strip()
    password = str(body.get("password", ""))
    confirmation = str(body.get("confirm_password", ""))
    if len(username) < 3 or len(username) > 80:
        raise HTTPException(status_code=422, detail="اسم المستخدم يجب أن يكون بين 3 و80 حرفاً.")
    if len(password) < 12 or len(password) > 256:
        raise HTTPException(status_code=422, detail="كلمة المرور يجب أن تكون 12 حرفاً على الأقل.")
    if password != confirmation:
        raise HTTPException(status_code=422, detail="تأكيد كلمة المرور غير مطابق.")
    hashed = password_hash(password)
    try:
        with connect() as conn:
            conn.execute(
                "INSERT INTO public.admin_account (id, username, password_hash) VALUES (1, %s, %s)",
                (username, hashed),
            )
    except psycopg.errors.UniqueViolation:
        raise HTTPException(status_code=409, detail="تم إنشاء حساب المدير مسبقاً.")
    response = JSONResponse({"authenticated": True, "username": username})
    response.set_cookie(
        COOKIE_NAME,
        session_token(username, int(time.time()) + SESSION_HOURS * 3600),
        max_age=SESSION_HOURS * 3600,
        httponly=True,
        secure=COOKIE_SECURE,
        samesite="strict",
        path="/",
    )
    return response


@app.get("/api/session")
def get_session(request: Request):
    username = authenticated_username(request)
    return {"authenticated": bool(username), "username": username}


@app.post("/api/login")
async def login(request: Request):
    check_same_origin(request)
    try:
        body = await request.json()
    except Exception:
        raise HTTPException(status_code=400, detail="بيانات تسجيل الدخول غير صالحة.")
    username = str(body.get("username", ""))[:200]
    password = str(body.get("password", ""))[:1000]
    client_ip = request.client.host if request.client else "unknown"
    bucket = f"{client_ip}:{username}"
    attempts = [stamp for stamp in LOGIN_FAILURES.get(bucket, []) if time.time() - stamp < 900]
    LOGIN_FAILURES[bucket] = attempts
    if len(attempts) >= 8:
        raise HTTPException(status_code=429, detail="محاولات كثيرة. حاول بعد 15 دقيقة.")
    account = get_admin_account()
    valid_user = bool(account and hmac.compare_digest(username.encode(), account["username"].encode()))
    valid_password = bool(account and verify_password(password, account["password_hash"]))
    if not (valid_user and valid_password):
        attempts.append(time.time())
        LOGIN_FAILURES[bucket] = attempts
        raise HTTPException(status_code=401, detail="اسم المستخدم أو كلمة المرور غير صحيحة.")
    LOGIN_FAILURES.pop(bucket, None)
    response = JSONResponse({"authenticated": True, "username": account["username"]})
    response.set_cookie(
        COOKIE_NAME,
        session_token(account["username"], int(time.time()) + SESSION_HOURS * 3600),
        max_age=SESSION_HOURS * 3600,
        httponly=True,
        secure=COOKIE_SECURE,
        samesite="strict",
        path="/",
    )
    return response


@app.post("/api/logout")
def logout(request: Request):
    check_same_origin(request)
    response = JSONResponse({"ok": True})
    response.delete_cookie(COOKIE_NAME, path="/", secure=COOKIE_SECURE, httponly=True, samesite="strict")
    return response


@app.get("/api/dashboard")
def dashboard():
    with connect() as conn:
        row = conn.execute(
            """
            SELECT count(*) AS total,
                   count(*) FILTER (WHERE employment_status = 'على رأس عمله') AS active,
                   count(*) FILTER (WHERE project IS NOT NULL AND btrim(project) <> '') AS project_staff,
                   count(DISTINCT work_center) AS centers
            FROM public.employees
            """
        ).fetchone()
        directorates = conn.execute(
            """SELECT coalesce(nullif(btrim(directorate), ''), 'غير محدد') AS name, count(*) AS count
               FROM public.employees GROUP BY 1 ORDER BY count DESC, name LIMIT 12"""
        ).fetchall()
        cadre_types = conn.execute(
            """SELECT coalesce(nullif(btrim(cadre_type), ''), 'غير محدد') AS name, count(*) AS count
               FROM public.employees GROUP BY 1 ORDER BY count DESC, name LIMIT 10"""
        ).fetchall()
    return {"totals": row, "directorates": directorates, "cadre_types": cadre_types}


FILTER_COLUMNS = {
    "work_center": "work_center",
    "directorate": "directorate",
    "department": "department",
    "job_title": "job_title",
    "cadre_type": "cadre_type",
    "employment_status": "employment_status",
}


@app.get("/api/employee-options")
def employee_options(work_center: str = "", directorate: str = "", department: str = ""):
    dependencies = {
        "work_center": [],
        "directorate": [("work_center", work_center)],
        "department": [("work_center", work_center), ("directorate", directorate)],
        "job_title": [("directorate", directorate), ("department", department)],
        "cadre_type": [],
        "employment_status": [],
    }
    result: dict[str, list[str]] = {}
    with connect() as conn:
        for key, column in FILTER_COLUMNS.items():
            where_parts: list[str] = []
            params: list[str] = []
            for filter_column, value in dependencies[key]:
                if value:
                    where_parts.append(f"{filter_column} = %s")
                    params.append(value)
            where_sql = (" WHERE " + " AND ".join(where_parts) + " AND") if where_parts else " WHERE"
            rows = conn.execute(
                f"SELECT DISTINCT {column} AS value FROM public.employees{where_sql} {column} IS NOT NULL AND btrim({column}) <> ''",
                params,
            ).fetchall()
            result[key] = sorted((str(row["value"]) for row in rows), key=str.casefold)
    return result


@app.get("/api/employees")
def list_employees(
    q: str = "",
    work_center: str = "",
    directorate: str = "",
    department: str = "",
    job_title: str = "",
    cadre_type: str = "",
    employment_status: str = "",
    page: int = 1,
    page_size: int = 20,
):
    page = max(1, page)
    page_size = min(100, max(10, page_size))
    clauses: list[str] = []
    params: list[Any] = []
    if q.strip():
        clauses.append("(employee_number ILIKE %s OR employee_name ILIKE %s OR job_title ILIKE %s)")
        term = f"%{q.strip()}%"
        params.extend([term, term, term])
    for column, value in (
        ("work_center", work_center),
        ("directorate", directorate),
        ("department", department),
        ("job_title", job_title),
        ("cadre_type", cadre_type),
        ("employment_status", employment_status),
    ):
        if value.strip():
            clauses.append(f"{column} = %s")
            params.append(value.strip())
    where_sql = (" WHERE " + " AND ".join(clauses)) if clauses else ""
    with connect() as conn:
        total = conn.execute(
            f"SELECT count(*) AS total FROM public.employees{where_sql}", params
        ).fetchone()["total"]
        items = conn.execute(
            f"""SELECT employee_number, employee_name, work_center, directorate, department,
                       job_title, cadre_type, employment_status, project, grade, is_frozen
                FROM public.employees{where_sql}
                ORDER BY directorate NULLS LAST, employee_name
                LIMIT %s OFFSET %s""",
            [*params, page_size, (page - 1) * page_size],
        ).fetchall()
    return {
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": (total + page_size - 1) // page_size,
    }



@app.get("/api/employees/{employee_number}")
def employee_details(employee_number: str):
    with connect() as conn:
        employee = conn.execute(
            """SELECT employee_number, employee_name, work_center, directorate, department,
                      job_title, cadre_type, employment_status, salary, next_grade_due_date,
                      coverage_percentage, project, grade, employee_code, is_frozen, frozen_at, frozen_by
               FROM public.employees WHERE employee_number = %s""",
            (employee_number,),
        ).fetchone()
    if not employee:
        raise HTTPException(status_code=404, detail="لم يتم العثور على الموظف.")
    return employee


EDITABLE_FIELDS = (
    "employee_number", "employee_name", "work_center", "directorate", "department",
    "job_title", "cadre_type", "employment_status", "salary", "next_grade_due_date",
    "coverage_percentage", "project", "grade", "employee_code",
)
REQUIRED_EMPLOYEE_FIELDS = {
    "employee_number", "employee_name", "work_center", "job_title",
    "cadre_type", "employment_status", "employee_code",
}


def validate_employee_update(payload: dict[str, Any]) -> tuple[Any, ...]:
    employee: dict[str, Any] = {}
    for key in EDITABLE_FIELDS:
        value = payload.get(key)
        try:
            if key in ("employee_number", "employee_code"):
                employee[key] = clean_identifier(value)
            elif key in ("salary", "coverage_percentage"):
                employee[key] = clean_decimal(value)
            elif key == "next_grade_due_date":
                employee[key] = clean_date(value)
            else:
                employee[key] = clean_text(value)
        except (ValueError, InvalidOperation, OverflowError):
            raise HTTPException(status_code=422, detail=f"قيمة غير صالحة في الحقل: {key}")
    missing = sorted(key for key in REQUIRED_EMPLOYEE_FIELDS if not employee.get(key))
    if missing:
        raise HTTPException(status_code=422, detail="أكمل الحقول الإلزامية قبل الحفظ.")
    return tuple(employee[key] for key in EDITABLE_FIELDS)


def validate_employee_settings(conn, values: tuple[Any, ...]):
    employee = dict(zip(EDITABLE_FIELDS, values))
    for field, table, label in [
        ("work_center", "ref_work_centers", "مركز العمل"),
        ("cadre_type", "ref_cadre_types", "نوع الكادر"),
        ("employment_status", "ref_employment_statuses", "الحالة"),
    ]:
        if not conn.execute(f"SELECT 1 FROM public.{table} WHERE name=%s", (employee[field],)).fetchone():
            raise HTTPException(status_code=422, detail=f"القيمة المختارة في «{label}» غير موجودة في الإعدادات.")
    directorate = employee.get("directorate")
    department = employee.get("department")
    if directorate:
        linked = conn.execute("""SELECT 1 FROM public.ref_center_directorates x JOIN public.ref_work_centers c ON c.id=x.center_id JOIN public.ref_directorates d ON d.id=x.directorate_id WHERE c.name=%s AND d.name=%s""", (employee["work_center"], directorate)).fetchone()
        if not linked:
            raise HTTPException(status_code=422, detail="الإدارة المختارة غير مرتبطة بمركز العمل المحدد.")
    if department:
        if not directorate:
            raise HTTPException(status_code=422, detail="اختر الإدارة قبل اختيار القسم.")
        linked = conn.execute("""SELECT 1 FROM public.ref_directorate_departments x JOIN public.ref_directorates d ON d.id=x.directorate_id JOIN public.ref_departments p ON p.id=x.department_id WHERE d.name=%s AND p.name=%s""", (directorate, department)).fetchone()
        if not linked:
            raise HTTPException(status_code=422, detail="القسم المختار غير مرتبط بالإدارة المحددة.")
    title = employee.get("job_title")
    if title:
        if not directorate or not department:
            raise HTTPException(status_code=422, detail="اختر الإدارة والقسم قبل اختيار المسمى الوظيفي.")
        linked = conn.execute("""SELECT 1 FROM public.ref_directorate_department_job_titles x JOIN public.ref_directorates d ON d.id=x.directorate_id JOIN public.ref_departments p ON p.id=x.department_id JOIN public.ref_job_titles j ON j.id=x.job_title_id WHERE d.name=%s AND p.name=%s AND j.name=%s""", (directorate, department, title)).fetchone()
        if not linked:
            raise HTTPException(status_code=422, detail="المسمى الوظيفي المختار غير مرتبط بالقسم والإدارة في الإعدادات.")

@app.put("/api/employees/{employee_number}")
async def update_employee(employee_number: str, request: Request):
    check_same_origin(request)
    try:
        payload = await request.json()
    except Exception:
        raise HTTPException(status_code=400, detail="بيانات التعديل غير صالحة.")
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="بيانات التعديل غير صالحة.")
    values = validate_employee_update(payload)
    try:
        with connect() as conn:
            current = conn.execute(
                "SELECT is_frozen FROM public.employees WHERE employee_number = %s FOR UPDATE",
                (employee_number,),
            ).fetchone()
            if not current:
                raise HTTPException(status_code=404, detail="لم يتم العثور على الموظف.")
            if current["is_frozen"]:
                raise HTTPException(status_code=423, detail="ملف الموظف مجمّد. أعد تفعيله قبل التعديل.")
            validate_employee_settings(conn, values)
            conn.execute(
                """UPDATE public.employees SET
                       employee_number=%s, employee_name=%s, work_center=%s, directorate=%s,
                       department=%s, job_title=%s, cadre_type=%s, employment_status=%s,
                       salary=%s, next_grade_due_date=%s, coverage_percentage=%s, project=%s,
                       grade=%s, employee_code=%s, updated_at=NOW()
                   WHERE employee_number=%s""",
                (*values, employee_number),
            )
    except HTTPException:
        raise
    except psycopg.errors.UniqueViolation:
        raise HTTPException(status_code=409, detail="رقم الموظف أو كوده مستخدم في سجل آخر.")
    except psycopg.errors.CheckViolation:
        raise HTTPException(status_code=422, detail="تحقق من أن الراتب ونسبة التغطية أرقام غير سالبة.")
    except psycopg.Error:
        raise HTTPException(status_code=500, detail="تعذر حفظ التعديلات.")
    return {"ok": True, "employee_number": values[0]}


@app.post("/api/employees/{employee_number}/freeze")
def freeze_employee(employee_number: str, request: Request):
    check_same_origin(request)
    username = authenticated_username(request) or "admin"
    with connect() as conn:
        row = conn.execute(
            """UPDATE public.employees
               SET is_frozen=TRUE, frozen_at=NOW(), frozen_by=%s, updated_at=NOW()
               WHERE employee_number=%s
               RETURNING employee_number""",
            (username, employee_number),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="لم يتم العثور على الموظف.")
    return {"ok": True, "is_frozen": True}


@app.post("/api/employees/{employee_number}/unfreeze")
def unfreeze_employee(employee_number: str, request: Request):
    check_same_origin(request)
    with connect() as conn:
        row = conn.execute(
            """UPDATE public.employees
               SET is_frozen=FALSE, frozen_at=NULL, frozen_by=NULL, updated_at=NOW()
               WHERE employee_number=%s
               RETURNING employee_number""",
            (employee_number,),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="لم يتم العثور على الموظف.")
    return {"ok": True, "is_frozen": False}


@app.get("/api/import-status")
def import_status():
    with connect() as conn:
        state = conn.execute(
            """SELECT source_filename, imported_rows, imported_at
               FROM public.employee_import_control WHERE id = 1"""
        ).fetchone()
    return {"imported": bool(state), "record": state}


EXCEL_FIELDS = {
    "رقم الموظف": "employee_number",
    "اسم الموظف": "employee_name",
    "مركز العمل": "work_center",
    "الإدارة": "directorate",
    "القسم": "department",
    "المسمى الوظيفي": "job_title",
    "نوع الكادر": "cadre_type",
    "الحالة": "employment_status",
    "الراتب": "salary",
    "استحقاق الدرجة التالية": "next_grade_due_date",
    "نسبة التغطية": "coverage_percentage",
    "المشروع": "project",
    "الدرجة الوظيفية": "grade",
    "كود الموظف": "employee_code",
}
REQUIRED_COLUMNS = {
    "رقم الموظف", "اسم الموظف", "مركز العمل", "المسمى الوظيفي",
    "نوع الكادر", "الحالة", "كود الموظف",
}
INSERT_SQL = """
INSERT INTO public.employees (
    employee_number, employee_name, work_center, directorate, department, job_title,
    cadre_type, employment_status, salary, next_grade_due_date, coverage_percentage,
    project, grade, employee_code
) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
"""


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def clean_identifier(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(value)


def clean_decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
    else:
        text = str(value).strip().replace("٬", "").replace(",", "").translate(str.maketrans("٠١٢٣٤٥٦٧٨٩٫", "0123456789."))
        text = text.replace("%", "").replace(" ", "")
        text = re.sub(r"[^0-9.\-]", "", text)
        if not text:
            return None
        try:
            number = Decimal(text)
        except InvalidOperation:
            raise ValueError("قيمة رقمية غير صالحة.")
    if not number.is_finite():
        raise ValueError("قيمة رقمية غير صالحة.")
    return number


def clean_date(value: Any) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return from_excel(value).date()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError("صيغة تاريخ غير صالحة.")


def parse_employee_workbook(content: bytes) -> list[tuple[Any, ...]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook["الموظفون"] if "الموظفون" in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        raise HTTPException(status_code=422, detail="ملف Excel فارغ.")
    headers = {clean_text(value): index for index, value in enumerate(header_row) if clean_text(value)}
    missing = sorted(REQUIRED_COLUMNS - set(headers))
    if missing:
        raise HTTPException(status_code=422, detail={"message": "أعمدة أساسية مفقودة.", "columns": missing})

    data: list[tuple[Any, ...]] = []
    seen_numbers: set[str] = set()
    seen_codes: set[str] = set()
    errors: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        source = {name: row[index] if index < len(row) else None for name, index in headers.items()}
        target: dict[str, Any] = {}
        for source_name, dest_name in EXCEL_FIELDS.items():
            value = source.get(source_name)
            try:
                if dest_name in ("employee_number", "employee_code"):
                    target[dest_name] = clean_identifier(value)
                elif dest_name in ("salary", "coverage_percentage"):
                    target[dest_name] = clean_decimal(value)
                elif dest_name == "next_grade_due_date":
                    target[dest_name] = clean_date(value)
                else:
                    target[dest_name] = clean_text(value)
            except (ValueError, InvalidOperation, OverflowError):
                errors.append({"row": row_number, "field": source_name})
        required_fields = ("employee_number", "employee_name", "work_center", "job_title", "cadre_type", "employment_status", "employee_code")
        missing_fields = [field for field in required_fields if not target.get(field)]
        if missing_fields:
            errors.append({"row": row_number, "field": "required"})
        number = target.get("employee_number")
        code = target.get("employee_code")
        if number and number in seen_numbers:
            errors.append({"row": row_number, "field": "duplicate_employee_number"})
        if code and code in seen_codes:
            errors.append({"row": row_number, "field": "duplicate_employee_code"})
        if number:
            seen_numbers.add(number)
        if code:
            seen_codes.add(code)
        if not missing_fields:
            data.append(tuple(target.get(EXCEL_FIELDS[header]) for header in EXCEL_FIELDS))
    workbook.close()
    if errors:
        raise HTTPException(
            status_code=422,
            detail={"message": "توجد أخطاء في الملف؛ لم يتم حفظ أي سجل.", "errors": errors[:50], "error_count": len(errors)},
        )
    if not data:
        raise HTTPException(status_code=422, detail="لم يتم العثور على سجلات موظفين صالحة في الملف.")
    return data


@app.post("/api/employees/import")
async def import_employees(request: Request, file: UploadFile = File(...)):
    check_same_origin(request)
    filename = Path(file.filename or "").name
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=415, detail="ارفع ملف Excel بصيغة .xlsx.")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="حجم الملف أكبر من 16 ميجابايت.")
    if not content:
        raise HTTPException(status_code=400, detail="الملف فارغ.")
    try:
        rows = parse_employee_workbook(content)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=422, detail="تعذر قراءة ملف Excel. تأكد أنه ملف .xlsx سليم.")
    file_hash = hashlib.sha256(content).hexdigest()
    try:
        with connect() as conn:
            conn.execute("SELECT pg_advisory_xact_lock(%s)", (7210042026,))
            if conn.execute("SELECT id FROM public.employee_import_control WHERE id = 1").fetchone():
                raise HTTPException(status_code=409, detail="تم استيراد ملف الموظفين مسبقاً؛ الاستيراد متاح مرة واحدة فقط.")
            with conn.cursor() as cursor:
                cursor.executemany(INSERT_SQL, rows)
            conn.execute(
                """INSERT INTO public.employee_import_control
                   (id, source_filename, sha256, imported_rows, imported_by)
                   VALUES (1, %s, %s, %s, %s)""",
                (filename, file_hash, len(rows), authenticated_username(request) or "admin"),
            )
    except HTTPException:
        raise
    except psycopg.errors.UniqueViolation:
        raise HTTPException(status_code=409, detail="يوجد رقم موظف أو كود موظف مستخدم مسبقاً؛ تم التراجع عن الاستيراد بالكامل.")
    except psycopg.Error:
        raise HTTPException(status_code=500, detail="تعذر حفظ البيانات؛ تم التراجع عن الاستيراد بالكامل.")
    return {"ok": True, "imported_rows": len(rows), "source_filename": filename}


@app.get("/")
def root_page(request: Request):
    if authenticated_username(request) is None:
        return RedirectResponse("/login.html", status_code=303)
    return RedirectResponse("/index.html", status_code=307)




REF_TABLES = {
    "work_centers": ("ref_work_centers", "work_center", "مركز العمل"),
    "directorates": ("ref_directorates", "directorate", "الدائرة / الإدارة"),
    "departments": ("ref_departments", "department", "القسم"),
    "job_titles": ("ref_job_titles", "job_title", "المسمى الوظيفي"),
    "employment_statuses": ("ref_employment_statuses", "employment_status", "الحالة"),
    "cadre_types": ("ref_cadre_types", "cadre_type", "نوع الكادر"),
}
REF_PARENTS = {
    "directorates": ("ref_center_directorates", "center_id", "directorate_id", "ref_work_centers", "work_center", "directorate"),
    "departments": ("ref_directorate_departments", "directorate_id", "department_id", "ref_directorates", "directorate", "department"),
}
REF_LINKS = {
    "center_directorates": ("ref_center_directorates", "center_id", "directorate_id", "ref_work_centers", "ref_directorates", "work_center", "directorate"),
    "directorate_departments": ("ref_directorate_departments", "directorate_id", "department_id", "ref_directorates", "ref_departments", "directorate", "department"),
    "department_job_titles": ("ref_department_job_titles", "department_id", "job_title_id", "ref_departments", "ref_job_titles", "department", "job_title"),
}


def valid_reference_category(category: str):
    item = REF_TABLES.get(category)
    if not item:
        raise HTTPException(status_code=404, detail="نوع القائمة غير معروف.")
    return item


@app.get("/api/settings/reference-data")
def reference_data():
    with connect() as conn:
        data = {
            key: conn.execute(f"SELECT id, name FROM public.{table} ORDER BY name").fetchall()
            for key, (table, _employee_column, _label) in REF_TABLES.items()
        }
        data["center_directorates"] = conn.execute(
            """SELECT x.center_id AS left_id,x.directorate_id AS right_id,
                      c.name AS left_name,d.name AS right_name
               FROM public.ref_center_directorates x
               JOIN public.ref_work_centers c ON c.id=x.center_id
               JOIN public.ref_directorates d ON d.id=x.directorate_id
               ORDER BY c.name,d.name"""
        ).fetchall()
        data["directorate_departments"] = conn.execute(
            """SELECT x.directorate_id AS left_id,x.department_id AS right_id,
                      d.name AS left_name,p.name AS right_name
               FROM public.ref_directorate_departments x
               JOIN public.ref_directorates d ON d.id=x.directorate_id
               JOIN public.ref_departments p ON p.id=x.department_id
               ORDER BY d.name,p.name"""
        ).fetchall()
        data["projects"] = conn.execute("SELECT id,name FROM public.projects ORDER BY name").fetchall()
        data["directorate_department_paths"] = conn.execute(
            """SELECT x.directorate_id,x.department_id,
                      x.directorate_id::text || ':' || x.department_id::text AS path_id,
                      d.name AS directorate_name,p.name AS department_name
               FROM public.ref_directorate_departments x
               JOIN public.ref_directorates d ON d.id=x.directorate_id
               JOIN public.ref_departments p ON p.id=x.department_id
               ORDER BY d.name,p.name"""
        ).fetchall()
        data["department_job_titles"] = conn.execute(
            """SELECT x.directorate_id,x.department_id,
                      x.directorate_id::text || ':' || x.department_id::text AS path_id,
                      x.job_title_id AS right_id,
                      d.name || ' — ' || p.name AS left_name,j.name AS right_name
               FROM public.ref_directorate_department_job_titles x
               JOIN public.ref_directorates d ON d.id=x.directorate_id
               JOIN public.ref_departments p ON p.id=x.department_id
               JOIN public.ref_job_titles j ON j.id=x.job_title_id
               ORDER BY d.name,p.name,j.name"""
        ).fetchall()
    return data



def validate_parent_ids(category: str, parent_ids: Any, conn):
    if category == "job_titles":
        if not isinstance(parent_ids, list):
            raise HTTPException(status_code=422, detail="حدد الأقسام المرتبطة بالمسمى.")
        try:
            paths = []
            for value in parent_ids:
                parts = str(value).split(":")
                if len(parts) != 2: raise ValueError
                paths.append((int(parts[0]), int(parts[1])))
            paths = sorted(set(paths))
        except (TypeError, ValueError):
            raise HTTPException(status_code=422, detail="قائمة ارتباطات الإدارة والقسم غير صالحة.")
        if not paths:
            raise HTTPException(status_code=422, detail="حدد قسمًا مرتبطًا واحدًا على الأقل.")
        found = conn.execute(
            "SELECT directorate_id,department_id FROM public.ref_directorate_departments WHERE (directorate_id,department_id) IN (SELECT * FROM unnest(%s::bigint[],%s::bigint[]))",
            ([x[0] for x in paths],[x[1] for x in paths]),
        ).fetchall()
        if len(found) != len(paths):
            raise HTTPException(status_code=422, detail="أحد مسارات الإدارة والقسم غير موجود.")
        return [f"{a}:{b}" for a,b in paths]
    if category not in REF_PARENTS:
        return []
    if not isinstance(parent_ids, list):
        raise HTTPException(status_code=422, detail="حدد العناصر المرتبطة من الهيكل التنظيمي.")
    try:
        ids = sorted({int(value) for value in parent_ids})
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="قائمة الارتباطات غير صالحة.")
    if not ids:
        raise HTTPException(status_code=422, detail="حدد ارتباطًا واحدًا على الأقل من الهيكل التنظيمي.")
    _link_table, parent_id_col, _member_id_col, parent_table, _parent_emp_col, _member_emp_col = REF_PARENTS[category]
    found = conn.execute(f"SELECT id FROM public.{parent_table} WHERE id=ANY(%s)", (ids,)).fetchall()
    if len(found) != len(ids):
        raise HTTPException(status_code=422, detail="تتضمن الارتباطات المختارة قيمة غير موجودة.")
    return ids


def replace_reference_parents(conn, category: str, item_id: int, parent_ids: list[int], current_name: str):
    if category == "job_titles":
        wanted = {(int(v.split(":")[0]),int(v.split(":")[1])) for v in parent_ids}
        current = conn.execute("SELECT directorate_id,department_id FROM public.ref_directorate_department_job_titles WHERE job_title_id=%s",(item_id,)).fetchall()
        old = {(int(r["directorate_id"]),int(r["department_id"])) for r in current}
        removed = old - wanted
        if removed:
            usage = conn.execute(
                """SELECT count(*) AS n FROM public.employees e
                   JOIN public.ref_directorates d ON d.name=e.directorate
                   JOIN public.ref_departments p ON p.name=e.department
                   WHERE e.job_title=%s AND (d.id,p.id) IN (SELECT * FROM unnest(%s::bigint[],%s::bigint[]))""",
                (current_name,[x[0] for x in removed],[x[1] for x in removed]),
            ).fetchone()["n"]
            if usage:
                raise HTTPException(status_code=409,detail="لا يمكن إزالة مسار إدارة/قسم مستخدم في سجل موظف.")
        conn.execute("DELETE FROM public.ref_directorate_department_job_titles WHERE job_title_id=%s",(item_id,))
        for directorate_id,department_id in wanted:
            conn.execute("INSERT INTO public.ref_directorate_department_job_titles(directorate_id,department_id,job_title_id) VALUES (%s,%s,%s) ON CONFLICT DO NOTHING",(directorate_id,department_id,item_id))
        conn.execute("DELETE FROM public.ref_department_job_titles WHERE job_title_id=%s",(item_id,))
        for department_id in sorted({x[1] for x in wanted}):
            conn.execute("INSERT INTO public.ref_department_job_titles(department_id,job_title_id) VALUES (%s,%s) ON CONFLICT DO NOTHING",(department_id,item_id))
        return
    if category not in REF_PARENTS:
        return
    link_table, parent_id_col, member_id_col, parent_table, parent_emp_col, member_emp_col = REF_PARENTS[category]
    current_links = conn.execute(
        f"SELECT {parent_id_col} AS parent_id FROM public.{link_table} WHERE {member_id_col}=%s",
        (item_id,),
    ).fetchall()
    selected = set(parent_ids)
    for link_row in current_links:
        parent_id = link_row["parent_id"]
        if parent_id in selected:
            continue
        parent = conn.execute(f"SELECT name FROM public.{parent_table} WHERE id=%s", (parent_id,)).fetchone()
        if parent:
            if category == "departments":
                protected = conn.execute(
                    "SELECT count(*) AS n FROM public.ref_directorate_department_job_titles WHERE directorate_id=%s AND department_id=%s",
                    (parent_id,item_id),
                ).fetchone()["n"]
                if protected:
                    raise HTTPException(status_code=409,detail="لا يمكن إزالة الإدارة لأن هناك مسميات وظيفية مرتبطة بهذا القسم ضمنها.")
            usage = conn.execute(
                f"SELECT count(*) AS n FROM public.employees WHERE {parent_emp_col}=%s AND {member_emp_col}=%s",
                (parent["name"], current_name),
            ).fetchone()["n"]
            if usage:
                raise HTTPException(
                    status_code=409,
                    detail=f"لا يمكن إزالة «{parent['name']}» من الهيكل؛ يوجد {REF_TABLES[category][2]} مستخدم بهذا الارتباط في سجل موظف.",
                )
    conn.execute(f"DELETE FROM public.{link_table} WHERE {member_id_col}=%s", (item_id,))
    for parent_id in parent_ids:
        conn.execute(
            f"INSERT INTO public.{link_table}({parent_id_col},{member_id_col}) VALUES (%s,%s) ON CONFLICT DO NOTHING",
            (parent_id,item_id),
        )


@app.post("/api/settings/reference-data/{category}")
async def create_reference_value(category: str, request: Request):
    table, _employee_column, label = valid_reference_category(category)
    check_same_origin(request)
    try:
        payload = await request.json()
    except Exception:
        raise HTTPException(status_code=400, detail="بيانات القائمة غير صالحة.")
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="بيانات القائمة غير صالحة.")
    name = clean_text(payload.get("name"))
    if not name:
        raise HTTPException(status_code=422, detail=f"أدخل {label}.")
    try:
        with connect() as conn:
            parent_ids = validate_parent_ids(category, payload.get("parent_ids", []), conn)
            row = conn.execute(
                f"INSERT INTO public.{table}(name) VALUES (%s) RETURNING id,name",
                (name,),
            ).fetchone()
            if category in REF_PARENTS or category == "job_titles":
                replace_reference_parents(conn,category,row["id"],parent_ids,name)
    except psycopg.errors.UniqueViolation:
        raise HTTPException(status_code=409, detail="هذا الاسم موجود بالفعل في القائمة.")
    return row


@app.patch("/api/settings/reference-data/{category}/{item_id}")
async def rename_reference_value(category: str, item_id: int, request: Request):
    table, employee_column, label = valid_reference_category(category)
    check_same_origin(request)
    try:
        payload = await request.json()
    except Exception:
        raise HTTPException(status_code=400, detail="بيانات التعديل غير صالحة.")
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="بيانات التعديل غير صالحة.")
    name = clean_text(payload.get("name"))
    if not name:
        raise HTTPException(status_code=422, detail=f"أدخل {label}.")
    with connect() as conn:
        current = conn.execute(f"SELECT id,name FROM public.{table} WHERE id=%s FOR UPDATE", (item_id,)).fetchone()
        if not current:
            raise HTTPException(status_code=404, detail="العنصر غير موجود.")
        duplicate = conn.execute(f"SELECT id FROM public.{table} WHERE name=%s AND id<>%s", (name,item_id)).fetchone()
        if duplicate:
            raise HTTPException(status_code=409, detail="يوجد عنصر آخر بالاسم نفسه.")
        parent_ids = validate_parent_ids(category,payload.get("parent_ids",[]),conn) if category in REF_PARENTS or category == "job_titles" else []
        if category in REF_PARENTS or category == "job_titles":
            replace_reference_parents(conn,category,item_id,parent_ids,current["name"])
        conn.execute(f"UPDATE public.{table} SET name=%s WHERE id=%s", (name,item_id))
        conn.execute(
            f"UPDATE public.employees SET {employee_column}=%s,updated_at=NOW() WHERE {employee_column}=%s",
            (name,current["name"]),
        )
    return {"id": item_id, "name": name}


@app.delete("/api/settings/reference-data/{category}/{item_id}")
def delete_reference_value(category: str, item_id: int, request: Request):
    table, employee_column, label = valid_reference_category(category)
    check_same_origin(request)
    with connect() as conn:
        current = conn.execute(f"SELECT id,name FROM public.{table} WHERE id=%s FOR UPDATE", (item_id,)).fetchone()
        if not current:
            raise HTTPException(status_code=404, detail="العنصر غير موجود.")
        relation_count = 0
        if category == "work_centers":
            relation_count = conn.execute("SELECT count(*) AS n FROM public.ref_center_directorates WHERE center_id=%s",(item_id,)).fetchone()["n"]
        elif category == "directorates":
            relation_count = conn.execute("SELECT (SELECT count(*) FROM public.ref_center_directorates WHERE directorate_id=%s)+(SELECT count(*) FROM public.ref_directorate_departments WHERE directorate_id=%s)+(SELECT count(*) FROM public.ref_directorate_department_job_titles WHERE directorate_id=%s) AS n",(item_id,item_id,item_id)).fetchone()["n"]
        elif category == "departments":
            relation_count = conn.execute("SELECT (SELECT count(*) FROM public.ref_directorate_departments WHERE department_id=%s)+(SELECT count(*) FROM public.ref_department_job_titles WHERE department_id=%s)+(SELECT count(*) FROM public.ref_directorate_department_job_titles WHERE department_id=%s) AS n",(item_id,item_id,item_id)).fetchone()["n"]
        elif category == "job_titles":
            relation_count = conn.execute("SELECT (SELECT count(*) FROM public.ref_department_job_titles WHERE job_title_id=%s)+(SELECT count(*) FROM public.ref_directorate_department_job_titles WHERE job_title_id=%s) AS n",(item_id,item_id)).fetchone()["n"]
        used = conn.execute(f"SELECT count(*) AS n FROM public.employees WHERE {employee_column}=%s",(current["name"],)).fetchone()["n"]
        if relation_count or used:
            raise HTTPException(status_code=409,detail=f"لا يمكن حذف {label} «{current['name']}» لأنه مرتبط بموظفين أو بعلاقات تنظيمية. أزل الارتباطات أولًا.")
        conn.execute(f"DELETE FROM public.{table} WHERE id=%s",(item_id,))
    return {"ok": True}


@app.post("/api/settings/reference-links/{relation}")
async def create_reference_link(relation: str, request: Request):
    config = REF_LINKS.get(relation)
    if not config:
        raise HTTPException(status_code=404, detail="نوع الارتباط غير معروف.")
    link_table,left_col,right_col,left_table,right_table,left_key,right_key=config
    check_same_origin(request)
    try:
        payload=await request.json()
        left_id=int(payload.get("left_id"))
        right_id=int(payload.get("right_id"))
    except Exception:
        raise HTTPException(status_code=422,detail="اختر طرفي الارتباط.")
    try:
        with connect() as conn:
            exists=conn.execute(
                f"SELECT (SELECT count(*) FROM public.{left_table} WHERE id=%s)+(SELECT count(*) FROM public.{right_table} WHERE id=%s) AS n",
                (left_id,right_id),
            ).fetchone()["n"]
            if exists!=2:
                raise HTTPException(status_code=404,detail="أحد عناصر الارتباط غير موجود.")
            conn.execute(
                f"INSERT INTO public.{link_table}({left_col},{right_col}) VALUES (%s,%s) ON CONFLICT DO NOTHING",
                (left_id,right_id),
            )
    except psycopg.Error:
        raise HTTPException(status_code=500,detail="تعذر حفظ الارتباط.")
    return {"ok":True}


@app.delete("/api/settings/reference-links/{relation}/{left_id}/{right_id}")
def delete_reference_link(relation: str,left_id: int,right_id: int,request: Request):
    config=REF_LINKS.get(relation)
    if not config:
        raise HTTPException(status_code=404,detail="نوع الارتباط غير معروف.")
    link_table,left_col,right_col,left_table,right_table,left_key,right_key=config
    check_same_origin(request)
    with connect() as conn:
        left=conn.execute(f"SELECT name FROM public.{left_table} WHERE id=%s",(left_id,)).fetchone()
        right=conn.execute(f"SELECT name FROM public.{right_table} WHERE id=%s",(right_id,)).fetchone()
        if not left or not right:
            raise HTTPException(status_code=404,detail="أحد عناصر الارتباط غير موجود.")
        employee_usage=conn.execute(
            f"SELECT count(*) AS n FROM public.employees WHERE {left_key}=%s AND {right_key}=%s",
            (left["name"],right["name"]),
        ).fetchone()["n"]
        if employee_usage:
            raise HTTPException(status_code=409,detail="لا يمكن حذف هذا الارتباط لأنه مستخدم في سجل موظف.")
        conn.execute(
            f"DELETE FROM public.{link_table} WHERE {left_col}=%s AND {right_col}=%s",
            (left_id,right_id),
        )
    return {"ok":True}



def normalize_reference_label(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower().replace("ـ", "")
    text = "".join(ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch))
    text = text.translate(str.maketrans({"أ":"ا","إ":"ا","آ":"ا","ى":"ي","ة":"ه","ؤ":"و","ئ":"ي"}))
    return re.sub(r"\s+", " ", text).strip()


def _planning_ref_maps(conn):
    tables = {
        "work_centers":"ref_work_centers",
        "directorates":"ref_directorates",
        "departments":"ref_departments",
        "job_titles":"ref_job_titles",
    }
    return {
        key: {normalize_reference_label(row["name"]): row for row in conn.execute(
            f"SELECT id,name FROM public.{table} ORDER BY name"
        ).fetchall()}
        for key, table in tables.items()
    }


def _planning_match(refs, category: str, value: Any, prefix: bool = False):
    key = normalize_reference_label(value)
    pool = refs.get(category, {})
    if key in pool:
        return pool[key]
    if prefix and key:
        for normalized, row in pool.items():
            if normalized.startswith(key + " ") or key.startswith(normalized + " "):
                return row
    return None


def _planning_source_rows(workbook, service_line: str, source_filename: str, refs):
    if service_line == "hospitals":
        ws = next((workbook[name] for name in workbook.sheetnames if normalize_reference_label(name) == "مستشفيات"), None)
        if ws is None:
            raise HTTPException(status_code=422, detail="ملف المستشفيات لا يحتوي على ورقة «مستشفيات».")
        blocks = []
        for col in range(4, min(ws.max_column, 22) + 1, 3):
            loc = ws.cell(3, col).value
            if loc and "مستشفى" in str(loc) and "مجموع" not in str(loc):
                center = _planning_match(refs, "work_centers", loc, prefix=True)
                blocks.append((col, str(loc).strip(), center, None, None))
    elif service_line == "primary_care":
        ws = next((workbook[name] for name in workbook.sheetnames if "رعايه" in normalize_reference_label(name)), None)
        if ws is None:
            raise HTTPException(status_code=422, detail="ملف الرعاية الأولية لا يحتوي على ورقة الرعاية.")
        center = _planning_match(refs, "work_centers", "الرعاية الصحية الأولية")
        directorate = _planning_match(refs, "directorates", "العيادات والخدمات العلاجية")
        blocks = []
        for col in range(4, ws.max_column + 1, 3):
            loc = ws.cell(3, col).value
            if not loc or "مجموع" in str(loc):
                continue
            location = re.sub(r"\s*[-–]\s*مستوى رابع\s*", "", str(loc)).strip()
            department = _planning_match(refs, "departments", location)
            blocks.append((col, location, center, directorate, department))
    else:
        ws = next((workbook[name] for name in workbook.sheetnames if "اسعاف" in normalize_reference_label(name)), None)
        if ws is None:
            raise HTTPException(status_code=422, detail="ملف المستشفيات لا يحتوي على ورقة دائرة الإسعاف والطوارئ.")
        center = _planning_match(refs, "work_centers", "الاسعاف والطواري -غزة", prefix=True)
        dir_alias = {
            "إدارة الدائرة":"ادارة الاسعاف", "خدمة المعبر":"ادارة الاسعاف",
            "مركز الشمال":"مركز جباليا", "مركز غزة":"مركز غزة",
            "مركز دير البلح":"مركز دير البلح", "مركز خانيونس":"مركز خان يونس",
            "مركز رفح":"مركز رفح",
        }
        dept_alias = {
            "إدارة الدائرة":"ادارة الاسعاف", "خدمة المعبر":"",
            "مركز الشمال":"اسعاف جباليا", "مركز غزة":"اسعاف غزة",
            "مركز دير البلح":"اسعاف دير البلح", "مركز خانيونس":"اسعاف خان يونس",
            "مركز رفح":"اسعاف رفح",
        }
        blocks = []
        for col in range(4, ws.max_column + 1, 3):
            loc = ws.cell(3, col).value
            if not loc or "مجموع" in str(loc):
                continue
            location = str(loc).strip()
            directorate = _planning_match(refs, "directorates", dir_alias.get(location, ""))
            department = _planning_match(refs, "departments", dept_alias.get(location, ""))
            blocks.append((col, location, center, directorate, department))

    records = []
    for col, location, center, mapped_dir, mapped_dept in blocks:
        for values in ws.iter_rows(min_row=5, values_only=True):
            if len(values) < 3 or not isinstance(values[2], str):
                continue
            source_title = values[2].strip()
            if not source_title or "مجموع" in source_title:
                continue
            raw_required = values[col - 1] if len(values) >= col else None
            try:
                required = int(float(raw_required or 0))
            except (TypeError, ValueError, OverflowError):
                continue
            if required <= 0:
                continue
            source_dir = str(values[0] or "").strip()
            source_dept = str(values[1] or "").strip()
            if service_line == "hospitals":
                directorate = _planning_match(refs, "directorates", source_dir)
                department = _planning_match(refs, "departments", source_dept)
                if not department and source_dept.startswith("أقسام "):
                    department = _planning_match(refs, "departments", "قسم " + source_dept[len("أقسام "):])
            elif service_line == "primary_care":
                directorate = mapped_dir
                department = mapped_dept
            else:
                directorate = mapped_dir
                department = mapped_dept
            job = _planning_match(refs, "job_titles", source_title)
            aliases = {
                "طبيب نساء وتوليد":"طبيب نساء وولادة",
                "أخصائي العلاج الطبيعي":"أخصائي علاج طبيعي",
            }
            if not job and source_title in aliases:
                job = _planning_match(refs, "job_titles", aliases[source_title])
            records.append({
                "service_line": service_line,
                "location_name": location,
                "work_center_id": center["id"] if center else None,
                "directorate_id": directorate["id"] if directorate else None,
                "department_id": department["id"] if department else None,
                "job_title_id": job["id"] if job else None,
                "source_directorate": source_dir,
                "source_department": source_dept,
                "source_job_title": source_title,
                "required_count": required,
                "source_file": source_filename,
            })
    return records


@app.get("/api/planning")
def list_workforce_requirements(service_line: str = "", location_name: str = "", q: str = ""):
    where = []
    params = []
    if service_line in {"hospitals", "primary_care", "emergency"}:
        where.append("p.service_line=%s"); params.append(service_line)
    if location_name.strip():
        where.append("p.location_name=%s"); params.append(location_name.strip())
    if q.strip():
        where.append("(p.source_directorate ILIKE %s OR p.source_department ILIKE %s OR p.source_job_title ILIKE %s OR p.location_name ILIKE %s OR coalesce(j.name,'') ILIKE %s)")
        term = f"%{q.strip()}%"; params.extend([term,term,term,term,term])
    where_sql = ("WHERE " + " AND ".join(where)) if where else ""
    with connect() as conn:
        rows = conn.execute(
            f"""WITH actual_staff AS (
                   SELECT work_center,directorate,department,job_title,
                          count(*) AS available,
                          jsonb_agg(jsonb_build_object(
                              'employee_number',employee_number,
                              'employee_name',employee_name,
                              'work_center',work_center,
                              'directorate',directorate,
                              'department',department,
                              'job_title',job_title,
                              'cadre_type',cadre_type,
                              'employment_status',employment_status,
                              'is_frozen',is_frozen
                          ) ORDER BY employee_name) AS employees
                   FROM public.employees
                   WHERE employment_status='على رأس عمله'
                   GROUP BY work_center,directorate,department,job_title
               )
               SELECT p.id,p.service_line,p.location_name,p.source_directorate,p.source_department,
                      p.source_job_title,p.required_count,p.source_file,p.updated_at,
                      p.work_center_id,c.name AS work_center,
                      p.directorate_id,d.name AS directorate,
                      p.department_id,dep.name AS department,
                      p.job_title_id,j.name AS mapped_job_title,
                      coalesce(s.available,0) AS available,
                      coalesce(s.employees,'[]'::jsonb) AS employees,
                      (p.work_center_id IS NOT NULL AND p.directorate_id IS NOT NULL
                       AND p.department_id IS NOT NULL AND p.job_title_id IS NOT NULL
                       AND EXISTS (SELECT 1 FROM public.ref_center_directorates cd
                                   WHERE cd.center_id=p.work_center_id AND cd.directorate_id=p.directorate_id)
                       AND EXISTS (SELECT 1 FROM public.ref_directorate_departments dd
                                   WHERE dd.directorate_id=p.directorate_id AND dd.department_id=p.department_id)
                       AND EXISTS (SELECT 1 FROM public.ref_directorate_department_job_titles jt
                                   WHERE jt.directorate_id=p.directorate_id AND jt.department_id=p.department_id AND jt.job_title_id=p.job_title_id)
                      ) AS mapping_complete
               FROM public.workforce_requirements p
               LEFT JOIN public.ref_work_centers c ON c.id=p.work_center_id
               LEFT JOIN public.ref_directorates d ON d.id=p.directorate_id
               LEFT JOIN public.ref_departments dep ON dep.id=p.department_id
               LEFT JOIN public.ref_job_titles j ON j.id=p.job_title_id
               LEFT JOIN actual_staff s ON s.work_center=c.name AND s.directorate=d.name
                       AND s.department=dep.name AND s.job_title=j.name
               {where_sql}
               ORDER BY p.service_line,p.location_name,p.source_directorate,p.source_department,p.source_job_title""",
            params,
        ).fetchall()
    for row in rows:
        row["shortage"] = int(row["required_count"]) - int(row["available"])
    return {"items": rows, "total": len(rows)}


def _validate_requirement_payload(conn, payload):
    try:
        work_center_id = int(payload.get("work_center_id"))
        directorate_id = int(payload.get("directorate_id"))
        department_id = int(payload.get("department_id"))
        job_title_id = int(payload.get("job_title_id"))
        required = int(payload.get("required_count"))
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="أكمل مركز العمل والإدارة والقسم والمسمى والاحتياج.")
    if required < 0:
        raise HTTPException(status_code=422, detail="الاحتياج لا يمكن أن يكون أقل من صفر.")
    path = conn.execute(
        """SELECT c.name AS center,d.name AS directorate,dep.name AS department,j.name AS job_title
           FROM public.ref_work_centers c
           JOIN public.ref_directorates d ON d.id=%s
           JOIN public.ref_departments dep ON dep.id=%s
           JOIN public.ref_job_titles j ON j.id=%s
           WHERE c.id=%s
             AND EXISTS (SELECT 1 FROM public.ref_center_directorates x WHERE x.center_id=c.id AND x.directorate_id=d.id)
             AND EXISTS (SELECT 1 FROM public.ref_directorate_departments x WHERE x.directorate_id=d.id AND x.department_id=dep.id)
             AND EXISTS (SELECT 1 FROM public.ref_directorate_department_job_titles x WHERE x.directorate_id=d.id AND x.department_id=dep.id AND x.job_title_id=j.id)""",
        (directorate_id,department_id,job_title_id,work_center_id),
    ).fetchone()
    if not path:
        raise HTTPException(status_code=422, detail="المسمى أو القسم غير مرتبط بالمسار التنظيمي المختار في الإعدادات.")
    return work_center_id,directorate_id,department_id,job_title_id,required,path


@app.post("/api/planning")
async def create_workforce_requirement(request: Request):
    check_same_origin(request)
    payload = await request.json()
    if not isinstance(payload,dict):
        raise HTTPException(status_code=400,detail="بيانات الاحتياج غير صالحة.")
    service_line = payload.get("service_line")
    if service_line not in {"hospitals","primary_care","emergency"}:
        raise HTTPException(status_code=422,detail="اختر نوع التخطيط.")
    location = clean_text(payload.get("location_name"))
    if not location:
        raise HTTPException(status_code=422,detail="أدخل اسم المنشأة أو الموقع.")
    with connect() as conn:
        center_id,directorate_id,department_id,job_id,required,path = _validate_requirement_payload(conn,payload)
        try:
            row=conn.execute(
                """INSERT INTO public.workforce_requirements(
                       service_line,location_name,work_center_id,directorate_id,department_id,job_title_id,
                       source_directorate,source_department,source_job_title,required_count,source_file
                   ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'إدخال يدوي')
                   RETURNING id""",
                (service_line,location,center_id,directorate_id,department_id,job_id,
                 path["directorate"],path["department"],path["job_title"],required),
            ).fetchone()
        except psycopg.errors.UniqueViolation:
            raise HTTPException(status_code=409,detail="يوجد احتياج مسجل بالمسار نفسه لهذا الموقع.")
    return {"id":row["id"],"ok":True}


@app.patch("/api/planning/{item_id}")
async def update_workforce_requirement(item_id: int, request: Request):
    check_same_origin(request)
    payload = await request.json()
    if not isinstance(payload,dict):
        raise HTTPException(status_code=400,detail="بيانات الاحتياج غير صالحة.")
    service_line=payload.get("service_line")
    if service_line not in {"hospitals","primary_care","emergency"}:
        raise HTTPException(status_code=422,detail="اختر نوع التخطيط.")
    location=clean_text(payload.get("location_name"))
    if not location:
        raise HTTPException(status_code=422,detail="أدخل اسم المنشأة أو الموقع.")
    with connect() as conn:
        current=conn.execute("SELECT id FROM public.workforce_requirements WHERE id=%s FOR UPDATE",(item_id,)).fetchone()
        if not current: raise HTTPException(status_code=404,detail="سجل الاحتياج غير موجود.")
        center_id,directorate_id,department_id,job_id,required,path = _validate_requirement_payload(conn,payload)
        try:
            conn.execute(
                """UPDATE public.workforce_requirements
                   SET service_line=%s,location_name=%s,work_center_id=%s,directorate_id=%s,department_id=%s,
                       job_title_id=%s,required_count=%s,updated_at=NOW()
                   WHERE id=%s""",
                (service_line,location,center_id,directorate_id,department_id,job_id,required,item_id),
            )
        except psycopg.errors.UniqueViolation:
            raise HTTPException(status_code=409,detail="يوجد احتياج آخر بالمسار نفسه لهذا الموقع.")
    return {"ok":True}


@app.delete("/api/planning/{item_id}")
def delete_workforce_requirement(item_id: int, request: Request):
    check_same_origin(request)
    with connect() as conn:
        cur=conn.execute("DELETE FROM public.workforce_requirements WHERE id=%s RETURNING id",(item_id,))
        if not cur.fetchone(): raise HTTPException(status_code=404,detail="سجل الاحتياج غير موجود.")
    return {"ok":True}


@app.post("/api/planning/import")
async def import_workforce_plans(request: Request, hospitals_file: UploadFile = File(...), primary_care_file: UploadFile = File(...)):
    check_same_origin(request)
    uploads=[]
    for upload in (hospitals_file,primary_care_file):
        if not upload.filename or not upload.filename.lower().endswith(".xlsx"):
            raise HTTPException(status_code=415,detail="ارفع ملفين بصيغة Excel .xlsx.")
        contents=await upload.read()
        if len(contents)>MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413,detail="حجم أحد الملفات أكبر من الحد المسموح.")
        uploads.append((upload.filename,contents))
    try:
        hospital_wb=load_workbook(io.BytesIO(uploads[0][1]),data_only=True,read_only=True)
        primary_wb=load_workbook(io.BytesIO(uploads[1][1]),data_only=True,read_only=True)
    except Exception:
        raise HTTPException(status_code=422,detail="تعذر قراءة ملفي Excel. تحقق من صحة الملفين.")
    with connect() as conn:
        refs=_planning_ref_maps(conn)
        records=_planning_source_rows(hospital_wb,"hospitals",uploads[0][0],refs)
        records+=_planning_source_rows(hospital_wb,"emergency",uploads[0][0],refs)
        records+=_planning_source_rows(primary_wb,"primary_care",uploads[1][0],refs)
        if not records:
            raise HTTPException(status_code=422,detail="لم أعثر على قيم احتياج أكبر من صفر في الملفين.")
        with conn.cursor() as cur:
            cur.executemany(
                """INSERT INTO public.workforce_requirements(
                       service_line,location_name,work_center_id,directorate_id,department_id,job_title_id,
                       source_directorate,source_department,source_job_title,required_count,source_file
                   ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                   ON CONFLICT(service_line,location_name,source_directorate,source_department,source_job_title)
                   DO UPDATE SET work_center_id=EXCLUDED.work_center_id,directorate_id=EXCLUDED.directorate_id,
                       department_id=EXCLUDED.department_id,job_title_id=EXCLUDED.job_title_id,
                       required_count=EXCLUDED.required_count,source_file=EXCLUDED.source_file,updated_at=NOW()""",
                [(r["service_line"],r["location_name"],r["work_center_id"],r["directorate_id"],r["department_id"],r["job_title_id"],
                  r["source_directorate"],r["source_department"],r["source_job_title"],r["required_count"],r["source_file"]) for r in records],
            )
    return {
        "imported_rows":len(records),
        "unmapped_titles":len({r["source_job_title"] for r in records if not r["job_title_id"]}),
        "unmapped_paths":sum(1 for r in records if not all([r["work_center_id"],r["directorate_id"],r["department_id"],r["job_title_id"]])),
    }



def _project_position_title_id(conn, title_label: str):
    normalized = normalize_reference_label(title_label)
    rows = conn.execute("SELECT id,name FROM public.ref_job_titles ORDER BY name").fetchall()
    for row in rows:
        if normalize_reference_label(row["name"]) == normalized:
            return row["id"], row["name"]
    return None, title_label


@app.get("/api/projects")
def list_projects():
    with connect() as conn:
        projects = conn.execute(
            "SELECT id,name,description,start_date,end_date,created_at,updated_at FROM public.projects ORDER BY name"
        ).fetchall()
        positions = conn.execute(
            """SELECT pp.id,pp.project_id,pp.job_title_id,pp.title_label,pp.planned_count,
                      pp.budget_amount,j.name AS mapped_title
               FROM public.project_positions pp
               LEFT JOIN public.ref_job_titles j ON j.id=pp.job_title_id
               ORDER BY pp.project_id,pp.title_label"""
        ).fetchall()
        staff = conn.execute(
            """SELECT employee_number,employee_name,project,job_title,work_center,directorate,
                      department,cadre_type,employment_status,is_frozen
               FROM public.employees
               WHERE NULLIF(BTRIM(project),'') IS NOT NULL
               ORDER BY project,job_title,employee_name"""
        ).fetchall()
    position_map = {}
    for pos in positions:
        pos["employees"] = []
        pos["current_count"] = 0
        pos["title"] = pos["mapped_title"] or pos["title_label"]
        position_map.setdefault(pos["project_id"], []).append(pos)
    project_by_name = {row["name"]: row for row in projects}
    for row in projects:
        row["positions"] = position_map.get(row["id"], [])
        row["employees"] = []
    for emp in staff:
        project = project_by_name.get(emp["project"])
        if not project:
            continue
        project["employees"].append(emp)
        title = normalize_reference_label(emp["job_title"])
        for pos in project["positions"]:
            if normalize_reference_label(pos["title"]) == title:
                pos["employees"].append(emp)
                pos["current_count"] += 1
                break
    for row in projects:
        row["assigned_count"] = len(row["employees"])
        row["position_count"] = len(row["positions"])
        row["total_planned"] = sum(int(p["planned_count"]) for p in row["positions"])
        row["total_budget"] = sum(int(p["planned_count"] or 0) * float(p["budget_amount"] or 0) for p in row["positions"])
        for pos in row["positions"]:
            pos["budget_amount"] = float(pos["budget_amount"] or 0)
    return {"items": projects, "total": len(projects)}


@app.post("/api/projects/discover")
async def discover_projects_from_employee_roster(request: Request, file: UploadFile = File(...)):
    check_same_origin(request)
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=415, detail="ارفع كشف الموظفين بصيغة Excel .xlsx.")
    content = await file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="حجم الملف أكبر من الحد المسموح.")
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(status_code=422, detail="تعذر قراءة ملف Excel. تحقق من صحة الملف.")
    sheet = None
    header = None
    required_headers = {
        "رقم الموظف": normalize_reference_label("رقم الموظف"),
        "المشروع": normalize_reference_label("المشروع"),
        "المسمى الوظيفي": normalize_reference_label("المسمى الوظيفي"),
    }
    for ws in workbook.worksheets:
        for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True), start=1):
            normalized = {normalize_reference_label(value): idx for idx, value in enumerate(row) if value is not None}
            if all(value in normalized for value in required_headers.values()):
                sheet, header = ws, normalized
                break
        if sheet:
            break
    if not sheet or not header:
        raise HTTPException(status_code=422, detail="لم أجد أعمدة رقم الموظف والمشروع والمسمى الوظيفي في الملف.")
    project_col = header[required_headers["المشروع"]]
    title_col = header[required_headers["المسمى الوظيفي"]]
    number_col = header[required_headers["رقم الموظف"]]
    records = []
    counts = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if len(row) <= max(project_col, title_col, number_col):
            continue
        raw_number, raw_project, raw_title = row[number_col], row[project_col], row[title_col]
        project_name = clean_text(raw_project)
        if not project_name:
            continue
        employee_number = clean_text(raw_number)
        if isinstance(raw_number, float) and raw_number.is_integer():
            employee_number = str(int(raw_number))
        title_label = clean_text(raw_title) or "مسمى غير محدد"
        records.append((employee_number, project_name, title_label))
        counts[(project_name, title_label)] = counts.get((project_name, title_label), 0) + 1
    if not records:
        raise HTTPException(status_code=422, detail="لم أعثر على موظفين مرتبطين بمشاريع في الكشف.")
    with connect() as conn:
        refs = _planning_ref_maps(conn)
        project_names = sorted({project for _, project, _ in records})
        for name in project_names:
            conn.execute(
                "INSERT INTO public.projects(name) VALUES (%s) ON CONFLICT(name) DO NOTHING",
                (name,),
            )
        for (project_name, title_label), assigned_count in counts.items():
            title = _planning_match(refs, "job_titles", title_label)
            conn.execute(
                """INSERT INTO public.project_positions(project_id,job_title_id,title_label,planned_count,budget_amount)
                   SELECT id,%s,%s,%s,0 FROM public.projects WHERE name=%s
                   ON CONFLICT(project_id,title_label) DO NOTHING""",
                (title["id"] if title else None, title_label, assigned_count, project_name),
            )
        employee_ids = [number for number, _, _ in records if number]
        matched = 0
        if employee_ids:
            existing = conn.execute(
                "SELECT employee_number FROM public.employees WHERE employee_number = ANY(%s)",
                (employee_ids,),
            ).fetchall()
            existing_ids = {str(row["employee_number"]) for row in existing}
            assignments = [(project, number) for number, project, _ in records if number and number in existing_ids]
            if assignments:
                with conn.cursor() as cur:
                    cur.executemany(
                        "UPDATE public.employees SET project=%s WHERE employee_number=%s",
                        assignments,
                    )
            matched = len(assignments)
    return {
        "projects_found": len(project_names),
        "positions_found": len(counts),
        "employees_matched": matched,
        "employees_not_found": len([1 for number, _, _ in records if number]) - matched,
    }


def _save_project_positions(conn, project_id: int, positions):
    if not isinstance(positions, list) or len(positions) > 200:
        raise HTTPException(status_code=422, detail="قائمة المسميات الوظيفية غير صالحة.")
    clean_positions = []
    seen = set()
    for item in positions:
        if not isinstance(item, dict):
            raise HTTPException(status_code=422, detail="أحد صفوف الوظائف غير صالح.")
        try:
            planned_count = int(item.get("planned_count", 0))
            budget_amount = Decimal(str(item.get("budget_amount", 0) or 0))
        except (ValueError, TypeError, InvalidOperation):
            raise HTTPException(status_code=422, detail="تأكد من صحة العدد والموازنة لكل مسمى.")
        if planned_count < 0 or budget_amount < 0:
            raise HTTPException(status_code=422, detail="العدد والموازنة لا يمكن أن يكونا سالبين.")
        raw_title_id = item.get("job_title_id")
        if raw_title_id:
            try:
                title_id = int(raw_title_id)
            except (ValueError, TypeError):
                raise HTTPException(status_code=422, detail="المسمى الوظيفي المختار غير صالح.")
            title = conn.execute("SELECT id,name FROM public.ref_job_titles WHERE id=%s", (title_id,)).fetchone()
            if not title:
                raise HTTPException(status_code=422, detail="المسمى الوظيفي غير موجود في الإعدادات.")
            title_label = title["name"]
        else:
            title_id = None
            title_label = clean_text(item.get("title_label"))
            if not title_label:
                raise HTTPException(status_code=422, detail="اختر المسمى الوظيفي لكل صف.")
        key = normalize_reference_label(title_label)
        if key in seen:
            raise HTTPException(status_code=422, detail="لا تكرر المسمى الوظيفي داخل المشروع.")
        seen.add(key)
        clean_positions.append((project_id, title_id, title_label, planned_count, budget_amount))
    conn.execute("DELETE FROM public.project_positions WHERE project_id=%s", (project_id,))
    if clean_positions:
        with conn.cursor() as cur:
            cur.executemany(
                """INSERT INTO public.project_positions(project_id,job_title_id,title_label,planned_count,budget_amount)
                   VALUES (%s,%s,%s,%s,%s)""",
                clean_positions,
            )


@app.post("/api/projects")
async def create_project(request: Request):
    check_same_origin(request)
    payload = await request.json()
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="بيانات المشروع غير صالحة.")
    name = clean_text(payload.get("name"))
    if not name:
        raise HTTPException(status_code=422, detail="أدخل اسم المشروع.")
    description = clean_text(payload.get("description"))
    with connect() as conn:
        try:
            row = conn.execute(
                """INSERT INTO public.projects(name,description,start_date,end_date)
                   VALUES (%s,%s,NULLIF(%s,'')::date,NULLIF(%s,'')::date) RETURNING id""",
                (name, description, payload.get("start_date") or "", payload.get("end_date") or ""),
            ).fetchone()
        except psycopg.errors.UniqueViolation:
            raise HTTPException(status_code=409, detail="اسم المشروع مسجل مسبقاً.")
        _save_project_positions(conn, row["id"], payload.get("positions", []))
    return {"ok": True, "id": row["id"]}


@app.put("/api/projects/{project_id}")
async def update_project(project_id: int, request: Request):
    check_same_origin(request)
    payload = await request.json()
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="بيانات المشروع غير صالحة.")
    name = clean_text(payload.get("name"))
    if not name:
        raise HTTPException(status_code=422, detail="أدخل اسم المشروع.")
    description = clean_text(payload.get("description"))
    with connect() as conn:
        current = conn.execute("SELECT id,name FROM public.projects WHERE id=%s FOR UPDATE", (project_id,)).fetchone()
        if not current:
            raise HTTPException(status_code=404, detail="المشروع غير موجود.")
        try:
            conn.execute(
                """UPDATE public.projects SET name=%s,description=%s,start_date=NULLIF(%s,'')::date,
                       end_date=NULLIF(%s,'')::date,updated_at=NOW() WHERE id=%s""",
                (name, description, payload.get("start_date") or "", payload.get("end_date") or "", project_id),
            )
        except psycopg.errors.UniqueViolation:
            raise HTTPException(status_code=409, detail="اسم المشروع مسجل مسبقاً.")
        if current["name"] != name:
            conn.execute("UPDATE public.employees SET project=%s WHERE project=%s", (name, current["name"]))
        _save_project_positions(conn, project_id, payload.get("positions", []))
    return {"ok": True}


@app.delete("/api/projects/{project_id}")
def delete_project(project_id: int, request: Request):
    check_same_origin(request)
    with connect() as conn:
        current = conn.execute("SELECT id,name FROM public.projects WHERE id=%s", (project_id,)).fetchone()
        if not current:
            raise HTTPException(status_code=404, detail="المشروع غير موجود.")
        linked = conn.execute("SELECT count(*) AS n FROM public.employees WHERE project=%s", (current["name"],)).fetchone()["n"]
        if linked:
            raise HTTPException(status_code=409, detail=f"لا يمكن حذف المشروع قبل نقل أو إزالة {linked} موظفاً مرتبطاً به.")
        conn.execute("DELETE FROM public.projects WHERE id=%s", (project_id,))
    return {"ok": True}



@app.get("/api/organization/structure")
def organization_structure(work_center: str = "", directorate: str = "", department: str = "", job_title: str = "", employment_status: str = "على رأس عمله"):
    conditions=[];params=[]
    for column,value in (("work_center",work_center),("directorate",directorate),("department",department),("job_title",job_title),("employment_status",employment_status)):
        if value.strip():
            conditions.append(f"{column}=%s");params.append(value.strip())
    where=("WHERE "+" AND ".join(conditions)) if conditions else ""
    with connect() as conn:
        rows=conn.execute(
            f"""SELECT work_center,directorate,department,job_title,count(*) AS employee_count,
                       jsonb_agg(jsonb_build_object('employee_number',employee_number,'employee_name',employee_name,
                           'cadre_type',cadre_type,'employment_status',employment_status,'is_frozen',is_frozen)
                           ORDER BY employee_name) AS employees
                FROM public.employees {where}
                GROUP BY work_center,directorate,department,job_title
                ORDER BY work_center,directorate,department,job_title""",
            params,
        ).fetchall()
    return {"items":rows,"total":len(rows)}

app.mount("/", StaticFiles(directory=PUBLIC_DIR, html=True), name="public")
